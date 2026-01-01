import json
import os
from pathlib import Path
from typing import Dict, Any

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openai import OpenAI

# ================== НАСТРОЙКИ ==================

BASE_URL = "https://shop.tastycoffee.ru/coffee"
PAGE_URL_TEMPLATE = BASE_URL + "?page={page}"

DATA_FILE = Path("tastycoffee_data_all.json")
XLS_FILE = Path("tastycoffee_report.xlsx")
TYPES_FILE = Path("coffee_types.json")

COFFEE_TYPES = [
    "Зерно кофе",
    "Дрип-пакеты",
    "Капсулы",
    "Готовый напиток",
    "Растворимый кофе",
    "Концентрат кофе",
]


# ================== ЗАГРУЗКА И ПАРСИНГ ==================


def fetch_page(page: int) -> str:
    """Скачиваем HTML нужной страницы каталога."""
    url = PAGE_URL_TEMPLATE.format(page=page)
    headers = {
        "User-Agent": "Mozilla/5.0 (compatible; TastyCoffeeWatcher/1.0)"
    }
    resp = requests.get(url, headers=headers, timeout=15)
    resp.raise_for_status()
    return resp.text


def parse_products(html: str) -> Dict[str, Dict[str, Any]]:
    """
    Парсим товары со страницы /coffee?page=N.

    Ориентируемся на вёрстку:
      - карточка товара: div.product.tc-tile-col.product-item
      - название: div.tc-tile__title div[itemprop='name'] a
      - цена: meta[itemprop='price'] внутри div.tc-tile__bottom
    """
    soup = BeautifulSoup(html, "lxml")

    products: Dict[str, Dict[str, Any]] = {}

    for card in soup.select("div.product.tc-tile-col.product-item"):
        # Название и ссылка
        title_link = card.select_one("div.tc-tile__title div[itemprop='name'] a")
        if not title_link:
            continue

        name = title_link.get_text(" ", strip=True)
        href = title_link.get("href") or ""

        if href.startswith("/"):
            url = "https://shop.tastycoffee.ru" + href
        else:
            url = href or BASE_URL

        # Цена
        price_meta = card.select_one("div.tc-tile__bottom meta[itemprop='price']")
        if not price_meta:
            price_text_el = card.select_one("div.tc-tile__bottom span.text-nowrap")
            if not price_text_el:
                continue
            price_text = price_text_el.get_text(strip=True)
            digits = "".join(ch for ch in price_text if ch.isdigit())
            if not digits:
                continue
            price = int(digits)
        else:
            try:
                price = int(price_meta.get("content"))
            except (TypeError, ValueError):
                continue

        products[url] = {
            "name": name,
            "price": price,
            "url": url,
        }

    return products


def fetch_all_products(max_pages: int = 50) -> Dict[str, Dict[str, Any]]:
    """
    Обходим /coffee?page=1..max_pages, пока есть новые товары.
    Возвращаем словарь всех уникальных товаров по URL.
    """
    all_products: Dict[str, Dict[str, Any]] = {}
    seen_urls = set()

    for page in range(1, max_pages + 1):
        try:
            html = fetch_page(page)
        except requests.RequestException as e:
            print(f"Ошибка загрузки страницы {page}: {e}")
            break

        products = parse_products(html)
        print(f"Страница {page}: найдено товаров {len(products)}")

        if not products:
            # пустая страница — дальше смысла нет
            break

        new_on_page = 0
        for url, data in products.items():
            if url not in seen_urls:
                seen_urls.add(url)
                all_products[url] = data
                new_on_page += 1

        if new_on_page == 0:
            print("Новых товаров на этой странице нет — останавливаемся.")
            break

    print(f"Всего уникальных товаров: {len(all_products)}")
    return all_products


# ================== JSON-СНАПШОТЫ ==================


def load_previous_data() -> Dict[str, Dict[str, Any]]:
    if not DATA_FILE.exists():
        return {}
    with DATA_FILE.open("r", encoding="utf-8") as f:
        return json.load(f)


def save_current_data(products: Dict[str, Dict[str, Any]]) -> None:
    with DATA_FILE.open("w", encoding="utf-8") as f:
        json.dump(products, f, ensure_ascii=False, indent=2)


# ================== КЭШ ТИПОВ КОФЕ ==================


def load_type_cache() -> Dict[str, str]:
    if not TYPES_FILE.exists():
        return {}
    with TYPES_FILE.open("r", encoding="utf-8") as f:
        return json.load(f)


def save_type_cache(cache: Dict[str, str]) -> None:
    with TYPES_FILE.open("w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=2)


# ================== КЛАССИФИКАЦИЯ ТИПА КОФЕ (OpenAI) ==================


def normalize_coffee_type(raw: str) -> str:
    """
    Приводим ответ модели к одному из допустимых значений COFFEE_TYPES.
    """
    raw_lower = (raw or "").strip().lower()

    # точное совпадение
    for t in COFFEE_TYPES:
        if raw_lower == t.lower():
            return t

    # частичное совпадение
    for t in COFFEE_TYPES:
        if t.lower() in raw_lower:
            return t

    # запасной вариант
    return "Зерно кофе"


def classify_coffee_type_with_openai(
    text: str,
    client: OpenAI,
    model: str = "gpt-4o-mini",
) -> str:
    """
    Классифицируем товар в один из заданных видов кофе.
    Используем Chat Completions API.
    """
    if not text:
        return "Зерно кофе"

    messages = [
        {
            "role": "system",
            "content": (
                "Ты помощник по классификации товаров. Тебе даётся текст, связанный с товаром кофе.\n"
                "Твоя задача — выбрать ОДИН тип кофе из списка:\n\n"
                "- Зерно кофе\n"
                "- Дрип-пакеты\n"
                "- Капсулы\n"
                "- Готовый напиток\n"
                "- Растворимый кофе\n"
                "- Концентрат кофе\n\n"
                "Отвечай строго ОДНОЙ строкой, БЕЗ дополнительных пояснений, ровно как в списке.\n"
                "Если не уверен, выбери наиболее подходящий вариант."
            ),
        },
        {
            "role": "user",
            "content": f"Текст товара:\n{text}",
        },
    ]

    response = client.chat.completions.create(
        model=model,
        messages=messages,
        temperature=0,
    )
    raw = response.choices[0].message.content or ""
    return normalize_coffee_type(raw)


def enrich_products_with_types(
    products: Dict[str, Dict[str, Any]],
    client: OpenAI,
) -> Dict[str, Dict[str, Any]]:
    """
    Для каждого товара проставляем поле 'coffee_type'.
    Используем кэш, чтобы не дёргать OpenAI повторно.
    """
    cache = load_type_cache()
    updated_cache = dict(cache)

    for url, data in products.items():
        if url in cache:
            coffee_type = cache[url]
        else:
            context = f"Название: {data['name']}\nURL: {url}"
            try:
                coffee_type = classify_coffee_type_with_openai(context, client)
            except Exception as e:
                print(f"Ошибка вызова OpenAI для {url}: {e}")
                coffee_type = "Зерно кофе"

            updated_cache[url] = coffee_type

        data["coffee_type"] = coffee_type

    save_type_cache(updated_cache)
    return products


# ================== СРАВНЕНИЕ ИЗМЕНЕНИЙ ==================


def compare_data(
    old: Dict[str, Dict[str, Any]],
    new: Dict[str, Dict[str, Any]],
) -> Dict[str, Any]:
    """
    Сравниваем старый и новый список товаров по URL.
    """
    old_keys = set(old.keys())
    new_keys = set(new.keys())

    new_items = [new[k] for k in sorted(new_keys - old_keys)]
    removed_items = [old[k] for k in sorted(old_keys - new_keys)]

    price_changes = []
    for key in sorted(old_keys & new_keys):
        old_price = old[key]["price"]
        new_price = new[key]["price"]
        if old_price != new_price:
            price_changes.append({
                "name": new[key]["name"],
                "url": new[key]["url"],
                "old_price": old_price,
                "new_price": new_price,
                "diff": new_price - old_price,
            })

    return {
        "new_items": new_items,
        "removed_items": removed_items,
        "price_changes": price_changes,
    }


# ================== ВЫГРУЗКА В EXCEL ==================


def export_to_excel(
    products: Dict[str, Dict[str, Any]],
    changes: Dict[str, Any],
    filename: Path = XLS_FILE,
) -> None:
    """
    Создаём Excel-файл со всеми текущими товарами.
    - Новые товары — полностью зелёная строка.
    - Изменившаяся цена — зелёная только ячейка с текущей ценой.
    """
    new_urls = {item["url"] for item in changes["new_items"]}
    changed_urls = {item["url"] for item in changes["price_changes"]}

    wb = Workbook()
    ws = wb.active
    ws.title = "Coffee"

    headers = [
        "Название",
        "Вид кофе",
        "Текущая цена, ₽",
        "URL",
        "Статус",
        "Старая цена, ₽",
        "Изменение, ₽",
    ]
    ws.append(headers)

    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

    old_data = load_previous_data()

    sorted_items = sorted(products.values(), key=lambda x: x["name"].lower())

    for product in sorted_items:
        url = product["url"]
        name = product["name"]
        price = product["price"]
        coffee_type = product.get("coffee_type", "")

        status = ""
        old_price = ""
        diff = ""

        if url in new_urls:
            status = "new"
        elif url in changed_urls:
            status = "price_changed"

        if url in old_data:
            old_price_val = old_data[url]["price"]
            old_price = old_price_val
            diff_val = price - old_price_val
            diff = diff_val

        row = [
            name,
            coffee_type,
            price,
            url,
            status,
            old_price,
            diff,
        ]
        ws.append(row)

        row_idx = ws.max_row

        if status == "new":
            # новый товар — красим всю строку
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = green_fill
        elif status == "price_changed":
            # изменённая цена — красим только ячейку с текущей ценой (колонка C, индекс 3)
            ws.cell(row=row_idx, column=3).fill = green_fill

    # Ширина колонок
    ws.column_dimensions["A"].width = 40  # Название
    ws.column_dimensions["B"].width = 18  # Вид кофе
    ws.column_dimensions["C"].width = 14  # Цена
    ws.column_dimensions["D"].width = 50  # URL
    ws.column_dimensions["E"].width = 16  # Статус
    ws.column_dimensions["F"].width = 16  # Старая цена
    ws.column_dimensions["G"].width = 16  # Изменение

    wb.save(filename)
    print(f"Excel-отчёт сохранён: {filename.resolve()}")


# ================== КОНСОЛЬНЫЙ ВЫВОД ==================


def print_changes(changes: Dict[str, Any]) -> None:
    new_items = changes["new_items"]
    removed_items = changes["removed_items"]
    price_changes = changes["price_changes"]

    if not new_items and not removed_items and not price_changes:
        print("Изменений ассортимента и цен не обнаружено.")
        return

    if new_items:
        print("\n=== НОВЫЕ ТОВАРЫ ===")
        for item in new_items:
            print(f"- {item['name']} — {item['price']} ₽ ({item['url']})")

    if removed_items:
        print("\n=== УДАЛЁННЫЕ ТОВАРЫ ===")
        for item in removed_items:
            print(f"- {item['name']} (ранее {item['price']} ₽) — {item['url']}")

    if price_changes:
        print("\n=== ИЗМЕНЕНИЕ ЦЕН ===")
        for item in price_changes:
            sign = "+" if item["diff"] > 0 else ""
            print(
                f"- {item['name']} — было {item['old_price']} ₽, "
                f"стало {item['new_price']} ₽ ({sign}{item['diff']} ₽) "
                f"({item['url']})"
            )


# ================== ТОЧКА ВХОДА ==================


def main():
    api_key = os.environ.get("OPENAI_API_KEY")
    if not api_key:
        raise RuntimeError(
            "Не задана переменная окружения OPENAI_API_KEY. "
            "Установи её локально или в Secrets на Replit."
        )

    client = OpenAI(api_key=api_key)

    current_products = fetch_all_products()

    if not current_products:
        print("Не удалось распарсить ни одного товара. Проверь, не изменилась ли вёрстка сайта.")
        return

    # Проставляем тип кофе через OpenAI (с кэшированием)
    current_products = enrich_products_with_types(current_products, client)

    previous_products = load_previous_data()

    if not previous_products:
        print("Первый запуск: сохраняю текущий список товаров со всех страниц.")
        save_current_data(current_products)

        empty_changes = {"new_items": [], "removed_items": [], "price_changes": []}
        export_to_excel(current_products, empty_changes)
        print(f"Сохранено товаров: {len(current_products)}")
        return

    changes = compare_data(previous_products, current_products)
    print_changes(changes)

    export_to_excel(current_products, changes)

    save_current_data(current_products)


if __name__ == "__main__":
    main()
